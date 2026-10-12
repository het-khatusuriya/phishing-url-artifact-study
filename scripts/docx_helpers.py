"""Shared python-docx helpers (bibliography, styles, markdown rendering) for the paper builders."""
import csv
import re
from datetime import date
from pathlib import Path

import bibtexparser
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "results" / "figures"
TAB = ROOT / "results" / "tables"
OUT = ROOT / "paper" / "Malicious_URL_Detection_Research_Paper.docx"

# ----------------------------------------------------------------------------- bibliography
BIB = {e["ID"]: e for e in bibtexparser.load(open(ROOT / "research/references.bib", encoding="utf-8")).entries}
CITE_ORDER = []


def cite_num(key):
    if key not in CITE_ORDER:
        CITE_ORDER.append(key)
    return CITE_ORDER.index(key) + 1


CITE_GROUP = re.compile(r"\[([a-z][a-z0-9_]*[†‡]?(?:\s*,\s*[a-z][a-z0-9_]*[†‡]?)*)\]")


def replace_citations(text):
    def rep(m):
        parts = [p.strip() for p in m.group(1).split(",")]
        keys = [p.rstrip("†‡") for p in parts]
        if not all(k in BIB for k in keys):
            return m.group(0)
        out = []
        for p, k in zip(parts, keys):
            out.append(f"{cite_num(k)}{p[len(k):]}")
        return "[" + ", ".join(out) + "]"
    return CITE_GROUP.sub(rep, text)


def clean_tex(s):
    s = re.sub(r"[{}]", "", s or "")
    s = s.replace("\\&", "&").replace("--", "–").replace("\\_", "_")
    return re.sub(r"\s+", " ", s).strip()


PARTICLES = {"le", "van", "von", "de", "der", "den", "della", "da", "di", "du", "la", "dos"}
CORPORATE = re.compile(r"(/|\(|Foundation|abuse\.ch|Spamhaus|OpenPhish|Amazon|Cisco|Kaggle|Internet|Common Crawl)", re.I)


def fmt_name(n):
    n = clean_tex(n)
    if "," in n:
        last, first = [x.strip() for x in n.split(",", 1)]
    else:
        toks = n.split()
        if len(toks) == 1:
            return n
        k = len(toks) - 1
        while k > 1 and toks[k - 1].lower() in PARTICLES:
            k -= 1
        last, first = " ".join(toks[k:]), " ".join(toks[:k])
    initials = " ".join("-".join(p[0] + "." for p in w.split("-") if p) for w in first.split() if w[0].isalpha())
    return f"{initials} {last}".strip()


def fmt_authors(raw):
    raw = raw or ""
    if CORPORATE.search(raw):
        return clean_tex(raw)
    names = [fmt_name(a) for a in re.split(r"\s+and\s+", raw) if a.strip()]
    if len(names) > 6:
        return names[0] + " et al."
    if len(names) == 2:
        return f"{names[0]} and {names[1]}"
    if len(names) > 2:
        return ", ".join(names[:-1]) + ", and " + names[-1]
    return names[0] if names else ""


def fmt_reference(e):
    authors = fmt_authors(e.get("author", ""))
    title = clean_tex(e.get("title", ""))
    venue = clean_tex(e.get("journal") or e.get("booktitle") or e.get("howpublished") or e.get("publisher") or "")
    if venue.lower().startswith("http"):
        venue = ""
    tail = []
    if e.get("journal"):
        tail.append(venue)
        for f, lab in (("volume", "vol."), ("number", "no."), ("pages", "pp.")):
            if e.get(f) and lab not in venue:
                tail.append(f"{lab} {clean_tex(e[f])}")
    elif e.get("booktitle"):
        tail.append("in " + venue)
        if e.get("pages") and "pp." not in venue:
            tail.append(f"pp. {clean_tex(e['pages'])}")
    elif venue:
        tail.append(venue)
    if e.get("eprint") and not e.get("doi") and "arxiv" not in venue.lower():
        tail.append(f"arXiv:{clean_tex(e['eprint'])}")
    if e.get("year") and clean_tex(e["year"]) not in venue:
        tail.append(clean_tex(e["year"]))
    s = (f"{authors}, " if authors else "") + f"“{title},” " + ", ".join(t for t in tail if t)
    if e.get("doi"):
        s += f", doi: {clean_tex(e['doi'])}"
    elif e.get("url"):
        s += f". [Online]. Available: {clean_tex(e['url'])}"
    return s.rstrip(". ") + "."


# ----------------------------------------------------------------------------- document helpers
doc = Document()
st = doc.styles
st["Normal"].font.name = "Times New Roman"
st["Normal"].font.size = Pt(11)
st["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st["Normal"].paragraph_format.space_after = Pt(6)
st["Normal"].paragraph_format.line_spacing = 1.15
for lvl, size in ((1, 14), (2, 12), (3, 11)):
    h = st[f"Heading {lvl}"]
    h.font.name = "Times New Roman"; h.font.size = Pt(size); h.font.bold = True
    h.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    rf = h._element.rPr.rFonts
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Times New Roman")
    h.paragraph_format.space_before = Pt(12 if lvl == 1 else 9); h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
cap = st["Caption"]; cap.font.name = "Times New Roman"; cap.font.size = Pt(9.5); cap.font.italic = False
cap.font.bold = False; cap.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for s_ in ("left_margin", "right_margin"):
    setattr(sec, s_, Cm(2.3))
sec.top_margin = sec.bottom_margin = Cm(2.2)
PORTRAIT_TEXT_WIDTH = 21.0 - 4.6
LANDSCAPE_TEXT_WIDTH = 29.7 - 3.0


def add_field(paragraph, instr):
    r = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "Right-click and choose Update Field."
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for el in (f1, it, f2, t, f3):
        r._r.append(el)


def footer_page_numbers(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(p.runs):
        r.text = ""
    add_field(p, "PAGE")


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*?\*|\[[^\]]+\]\([^)]+\))")


def add_inline(p, text, size=None, base_bold=False, base_italic=False):
    text = replace_citations(text)
    for tok in INLINE.split(text):
        if not tok:
            continue
        bold, italic, mono = base_bold, base_italic, False
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            tok, bold = tok[2:-2], True
        elif tok.startswith("`") and tok.endswith("`"):
            tok, mono = tok[1:-1], True
        elif re.fullmatch(r"\[[^\]]+\]\([^)]+\)", tok):
            m = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", tok); tok = f"{m.group(1)} ({m.group(2)})"
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, italic = tok[1:-1], True
        r = p.add_run(tok)
        r.bold, r.italic = bold, italic
        if mono:
            r.font.name = "Consolas"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt((size or 11) - 1.5)
        elif size:
            r.font.size = Pt(size)
    return p


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_fill)
    tcPr.append(sh)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)


def add_table(header, rows, width_cm, font=8.0, col_weights=None):
    n = len(header)
    t = doc.add_table(rows=1, cols=n)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    def _w(i):
        head = max((len(x) for x in str(header[i]).split()), default=4)
        body = max([len(str(r[i])) for r in rows[:300]] + [0])
        return max(6, head + 1, min(38, body))
    w = col_weights or [_w(i) for i in range(n)]
    tot = sum(w); widths = [Cm(width_cm * x / tot) for x in w]
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.width = widths[i]; shade(c, "D9E2F3")
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        add_inline(p, str(h), size=font, base_bold=True)
    set_repeat_header(t.rows[0])
    for r in rows:
        cells = t.add_row().cells
        for i in range(n):
            cells[i].width = widths[i]
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            val = "" if r[i] is None or (isinstance(r[i], float) and pd.isna(r[i])) else str(r[i])
            for k, part in enumerate(val.split(chr(10))):
                if k:
                    p = cells[i].add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
                add_inline(p, part, size=font)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


FIG_N = [0]
TAB_N = [0]
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV", "XV", "XVI",
         "XVII", "XVIII", "XIX", "XX", "XXI", "XXII", "XXIII", "XXIV", "XXV", "XXVI", "XXVII", "XXVIII", "XXIX", "XXX"]


def add_figure(path, caption, width_cm=15.5):
    path = Path(path)
    if not path.exists():
        p = doc.add_paragraph(); add_inline(p, f"[Figure file missing: {path.name} — pending]", base_italic=True); return
    FIG_N[0] += 1
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    from PIL import Image
    with Image.open(path) as im:
        aspect = im.height / im.width
    if aspect > 0.75:
        width_cm = min(width_cm, 9.5)
    elif aspect > 0.55:
        width_cm = min(width_cm, 13.5)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    c = doc.add_paragraph(style="Caption"); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Fig. {FIG_N[0]}. "); r.bold = True
    add_inline(c, caption, size=9.5)


def table_caption(caption, label=None):
    TAB_N[0] += 1
    c = doc.add_paragraph(style="Caption"); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.keep_with_next = True
    r = c.add_run(f"TABLE {label or ROMAN[TAB_N[0] - 1]}. "); r.bold = True
    add_inline(c, caption, size=9.5)


def new_section(landscape):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    if landscape:
        s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(1.5); s.top_margin = s.bottom_margin = Cm(1.6)
    else:
        s.orientation = WD_ORIENT.PORTRAIT; s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.3); s.top_margin = s.bottom_margin = Cm(2.2)
    footer_page_numbers(s)
    return s


# ----------------------------------------------------------------------------- markdown renderer
MARKER = re.compile(r"^<<(FIG|TABLE|CAPTION|MDTABLECAP)\|(.*)>>$")


def parse_md_table(lines):
    rows = []
    for l in lines:
        l = l.strip()
        if re.fullmatch(r"\|?\s*:?-{2,}.*", l):
            continue
        cells = [c.strip() for c in l.strip("|").split("|")]
        rows.append(cells)
    if not rows:
        return None, []
    n = max(len(r) for r in rows)
    rows = [r + [""] * (n - len(r)) for r in rows]
    return rows[0], rows[1:]


def render_md(text, heading_offset=0, width_cm=PORTRAIT_TEXT_WIDTH, table_font=8.0, skip_title=False, handlers=None):
    lines = text.splitlines()
    i, pending_caption = 0, None
    while i < len(lines):
        line = lines[i].rstrip()
        m = MARKER.match(line.strip())
        if m:
            kind, arg = m.groups()
            if kind == "MDTABLECAP":
                pending_caption = arg
            elif handlers and kind in handlers:
                handlers[kind](arg)
            i += 1; continue
        if line.startswith("```"):
            j = i + 1; block = []
            while j < len(lines) and not lines[j].startswith("```"):
                block.append(lines[j]); j += 1
            for b in block:
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
                r = p.add_run(b if b else " "); r.font.name = "Consolas"; r.font.size = Pt(8)
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            i = j + 1; continue
        if line.startswith("|"):
            j = i; tl = []
            while j < len(lines) and lines[j].strip().startswith("|"):
                tl.append(lines[j]); j += 1
            header, rows = parse_md_table(tl)
            if header:
                if pending_caption:
                    table_caption(pending_caption); pending_caption = None
                add_table(header, rows, width_cm, font=table_font)
            i = j; continue
        hm = re.match(r"^(#{1,4})\s+(.*)", line)
        if hm:
            lvl = len(hm.group(1)) + heading_offset
            if skip_title and len(hm.group(1)) == 1:
                i += 1; continue
            doc.add_heading(replace_citations(re.sub(r"\*\*|`", "", hm.group(2))), level=min(max(lvl, 1), 4))
            i += 1; continue
        if re.match(r"^\s*[-*]\s+", line):
            indent = len(line) - len(line.lstrip())
            p = doc.add_paragraph(style="List Bullet" if indent < 2 else "List Bullet 2")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, re.sub(r"^\s*[-*]\s+", "", line)); i += 1; continue
        if re.match(r"^\s*\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number"); p.paragraph_format.space_after = Pt(2)
            add_inline(p, re.sub(r"^\s*\d+\.\s+", "", line)); i += 1; continue
        if line.startswith(">"):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8)
            add_inline(p, line.lstrip("> "), base_italic=True); i += 1; continue
        if line.strip() in ("---", "***"):
            i += 1; continue
        if not line.strip():
            i += 1; continue
        # merge consecutive text lines into one paragraph
        j = i; buf = [line.strip()]
        while j + 1 < len(lines) and lines[j + 1].strip() and not re.match(r"^(#|\||```|\s*[-*]\s|\s*\d+\.\s|>|<<)", lines[j + 1]):
            j += 1; buf.append(lines[j].strip())
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, " ".join(buf))
        i = j + 1


