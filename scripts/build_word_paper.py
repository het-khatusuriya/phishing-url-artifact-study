"""Build the complete research paper as a Word document.

Output: paper/Malicious_URL_Detection_Research_Paper.docx

Contents (all generated from files in the repository; nothing typed by hand here
except captions and section glue):
  title page, table of contents, the manuscript (paper/manuscript.md) with
  figures and tables placed in the text, numbered IEEE-style references built
  from research/references.bib, and appendices:
    A literature survey table (all verified sources)          research/LITERATURE_SURVEY_TABLE.md + paper_metadata.csv
    B full literature review                                   research/LITERATURE_REVIEW.md
    C research gaps and hypotheses                             research/RESEARCH_GAPS.md
    D repository audit                                         research/REPOSITORY_AUDIT.md
    E dataset audit                                            research/DATASET_AUDIT.md
    F experiment protocol                                      research/EXPERIMENT_PROTOCOL.md
    G complete result tables T1-T10                            results/tables/paper_T*.csv
    H additional figures (incl. the original project's outputs) results/figures/*.png
    I decision log                                             research/DECISIONS.md
    J reproducibility guide                                    README.md
Citations [citekey] are converted to numbers in order of first appearance.

Usage: python scripts/build_word_paper.py
"""
import csv
import re
from datetime import date
from pathlib import Path

import bibtexparser
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "results" / "figures"
TAB = ROOT / "results" / "tables"
OUT = ROOT / "paper" / "Malicious_URL_Detection_Research_Paper.docx"

# ----------------------------------------------------------------------------- bibliography
BIB = {e["ID"]: e for e in bibtexparser.load(open(ROOT / "research/references.bib", encoding="utf-8")).entries}
CITE_ORDER = []


def cite_num(key):
    if key not in CITE_ORDER:
        CITE_ORDER.append(key)
    return CITE_ORDER.index(key) + 1


CITE_GROUP = re.compile(r"\[([a-z][a-z0-9_]*[†‡]?(?:\s*,\s*[a-z][a-z0-9_]*[†‡]?)*)\]")


def replace_citations(text):
    def rep(m):
        parts = [p.strip() for p in m.group(1).split(",")]
        keys = [p.rstrip("†‡") for p in parts]
        if not all(k in BIB for k in keys):
            return m.group(0)
        out = []
        for p, k in zip(parts, keys):
            out.append(f"{cite_num(k)}{p[len(k):]}")
        return "[" + ", ".join(out) + "]"
    return CITE_GROUP.sub(rep, text)


def clean_tex(s):
    s = re.sub(r"[{}]", "", s or "")
    s = s.replace("\\&", "&").replace("--", "–").replace("\\_", "_")
    return re.sub(r"\s+", " ", s).strip()


def fmt_name(n):
    n = clean_tex(n)
    if "," in n:
        last, first = [x.strip() for x in n.split(",", 1)]
    else:
        toks = n.split()
        if len(toks) == 1:
            return n
        last, first = toks[-1], " ".join(toks[:-1])
    initials = " ".join("-".join(p[0] + "." for p in w.split("-") if p) for w in first.split() if w[0].isalpha())
    return f"{initials} {last}".strip()


def fmt_reference(e):
    authors_raw = e.get("author", "")
    if authors_raw.strip().startswith("{") or " and " not in authors_raw and "," not in authors_raw and len(authors_raw.split()) > 4:
        authors = clean_tex(authors_raw)
    else:
        names = [fmt_name(a) for a in re.split(r"\s+and\s+", authors_raw) if a.strip()]
        if len(names) > 6:
            authors = names[0] + " et al."
        elif len(names) > 1:
            authors = ", ".join(names[:-1]) + ", and " + names[-1]
        else:
            authors = names[0] if names else ""
    title = clean_tex(e.get("title", ""))
    venue = clean_tex(e.get("journal") or e.get("booktitle") or e.get("howpublished") or e.get("publisher") or "")
    parts = [f"{authors}," if authors else "", f"“{title},”"]
    tail = []
    if e.get("journal"):
        tail.append(venue)
        if e.get("volume"):
            tail.append(f"vol. {clean_tex(e['volume'])}")
        if e.get("number"):
            tail.append(f"no. {clean_tex(e['number'])}")
        if e.get("pages"):
            tail.append(f"pp. {clean_tex(e['pages'])}")
    elif e.get("booktitle"):
        tail.append("in " + venue)
        if e.get("pages"):
            tail.append(f"pp. {clean_tex(e['pages'])}")
    elif venue:
        tail.append(venue)
    if e.get("eprint") and "arxiv" not in venue.lower():
        tail.append(f"arXiv:{clean_tex(e['eprint'])}")
    if e.get("year"):
        tail.append(clean_tex(e["year"]))
    s = " ".join(p for p in parts if p) + " " + ", ".join(t for t in tail if t)
    if e.get("doi"):
        s += f", doi: {clean_tex(e['doi'])}"
    elif e.get("url"):
        s += f". [Online]. Available: {clean_tex(e['url'])}"
    return s.rstrip(".") + "."


# ----------------------------------------------------------------------------- document helpers
doc = Document()
st = doc.styles
st["Normal"].font.name = "Times New Roman"
st["Normal"].font.size = Pt(11)
st["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st["Normal"].paragraph_format.space_after = Pt(6)
st["Normal"].paragraph_format.line_spacing = 1.15
for lvl, size in ((1, 14), (2, 12), (3, 11)):
    h = st[f"Heading {lvl}"]
    h.font.name = "Times New Roman"; h.font.size = Pt(size); h.font.bold = True
    h.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    rf = h._element.rPr.rFonts
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Times New Roman")
    h.paragraph_format.space_before = Pt(12 if lvl == 1 else 9); h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
cap = st["Caption"]; cap.font.name = "Times New Roman"; cap.font.size = Pt(9.5); cap.font.italic = False
cap.font.bold = False; cap.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for s_ in ("left_margin", "right_margin"):
    setattr(sec, s_, Cm(2.3))
sec.top_margin = sec.bottom_margin = Cm(2.2)
PORTRAIT_TEXT_WIDTH = 21.0 - 4.6
LANDSCAPE_TEXT_WIDTH = 29.7 - 3.0


def add_field(paragraph, instr):
    r = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "Right-click and choose Update Field."
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for el in (f1, it, f2, t, f3):
        r._r.append(el)


def footer_page_numbers(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(p.runs):
        r.text = ""
    add_field(p, "PAGE")


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*?\*|\[[^\]]+\]\([^)]+\))")


def add_inline(p, text, size=None, base_bold=False, base_italic=False):
    text = replace_citations(text)
    for tok in INLINE.split(text):
        if not tok:
            continue
        bold, italic, mono = base_bold, base_italic, False
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            tok, bold = tok[2:-2], True
        elif tok.startswith("`") and tok.endswith("`"):
            tok, mono = tok[1:-1], True
        elif re.fullmatch(r"\[[^\]]+\]\([^)]+\)", tok):
            m = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", tok); tok = f"{m.group(1)} ({m.group(2)})"
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, italic = tok[1:-1], True
        r = p.add_run(tok)
        r.bold, r.italic = bold, italic
        if mono:
            r.font.name = "Consolas"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt((size or 11) - 1.5)
        elif size:
            r.font.size = Pt(size)
    return p


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_fill)
    tcPr.append(sh)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)


def add_table(header, rows, width_cm, font=8.0, col_weights=None):
    n = len(header)
    t = doc.add_table(rows=1, cols=n)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    def _w(i):
        head = max((len(x) for x in str(header[i]).split()), default=4)
        body = max([len(str(r[i])) for r in rows[:300]] + [0])
        return max(6, head + 1, min(38, body))
    w = col_weights or [_w(i) for i in range(n)]
    tot = sum(w); widths = [Cm(width_cm * x / tot) for x in w]
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.width = widths[i]; shade(c, "D9E2F3")
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        add_inline(p, str(h), size=font, base_bold=True)
    set_repeat_header(t.rows[0])
    for r in rows:
        cells = t.add_row().cells
        for i in range(n):
            cells[i].width = widths[i]
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            add_inline(p, "" if r[i] is None or (isinstance(r[i], float) and pd.isna(r[i])) else str(r[i]), size=font)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


FIG_N = [0]
TAB_N = [0]
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV", "XV", "XVI",
         "XVII", "XVIII", "XIX", "XX", "XXI", "XXII", "XXIII", "XXIV", "XXV", "XXVI", "XXVII", "XXVIII", "XXIX", "XXX"]


def add_figure(path, caption, width_cm=15.5):
    path = Path(path)
    if not path.exists():
        p = doc.add_paragraph(); add_inline(p, f"[Figure file missing: {path.name} — pending]", base_italic=True); return
    FIG_N[0] += 1
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    from PIL import Image
    with Image.open(path) as im:
        aspect = im.height / im.width
    if aspect > 0.75:
        width_cm = min(width_cm, 9.5)
    elif aspect > 0.55:
        width_cm = min(width_cm, 13.5)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    c = doc.add_paragraph(style="Caption"); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Fig. {FIG_N[0]}. "); r.bold = True
    add_inline(c, caption, size=9.5)


def table_caption(caption, label=None):
    TAB_N[0] += 1
    c = doc.add_paragraph(style="Caption"); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.keep_with_next = True
    r = c.add_run(f"TABLE {label or ROMAN[TAB_N[0] - 1]}. "); r.bold = True
    add_inline(c, caption, size=9.5)


def new_section(landscape):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    if landscape:
        s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(1.5); s.top_margin = s.bottom_margin = Cm(1.6)
    else:
        s.orientation = WD_ORIENT.PORTRAIT; s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.3); s.top_margin = s.bottom_margin = Cm(2.2)
    footer_page_numbers(s)
    return s


# ----------------------------------------------------------------------------- markdown renderer
MARKER = re.compile(r"^<<(FIG|TABLE|CAPTION|MDTABLECAP)\|(.*)>>$")


def parse_md_table(lines):
    rows = []
    for l in lines:
        l = l.strip()
        if re.fullmatch(r"\|?\s*:?-{2,}.*", l):
            continue
        cells = [c.strip() for c in l.strip("|").split("|")]
        rows.append(cells)
    if not rows:
        return None, []
    n = max(len(r) for r in rows)
    rows = [r + [""] * (n - len(r)) for r in rows]
    return rows[0], rows[1:]


def render_md(text, heading_offset=0, width_cm=PORTRAIT_TEXT_WIDTH, table_font=8.0, skip_title=False, handlers=None):
    lines = text.splitlines()
    i, pending_caption = 0, None
    while i < len(lines):
        line = lines[i].rstrip()
        m = MARKER.match(line.strip())
        if m:
            kind, arg = m.groups()
            if kind == "MDTABLECAP":
                pending_caption = arg
            elif handlers and kind in handlers:
                handlers[kind](arg)
            i += 1; continue
        if line.startswith("```"):
            j = i + 1; block = []
            while j < len(lines) and not lines[j].startswith("```"):
                block.append(lines[j]); j += 1
            for b in block:
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
                r = p.add_run(b if b else " "); r.font.name = "Consolas"; r.font.size = Pt(8)
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            i = j + 1; continue
        if line.startswith("|"):
            j = i; tl = []
            while j < len(lines) and lines[j].strip().startswith("|"):
                tl.append(lines[j]); j += 1
            header, rows = parse_md_table(tl)
            if header:
                if pending_caption:
                    table_caption(pending_caption); pending_caption = None
                add_table(header, rows, width_cm, font=table_font)
            i = j; continue
        hm = re.match(r"^(#{1,4})\s+(.*)", line)
        if hm:
            lvl = len(hm.group(1)) + heading_offset
            if skip_title and len(hm.group(1)) == 1:
                i += 1; continue
            doc.add_heading(replace_citations(re.sub(r"\*\*|`", "", hm.group(2))), level=min(max(lvl, 1), 4))
            i += 1; continue
        if re.match(r"^\s*[-*]\s+", line):
            indent = len(line) - len(line.lstrip())
            p = doc.add_paragraph(style="List Bullet" if indent < 2 else "List Bullet 2")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, re.sub(r"^\s*[-*]\s+", "", line)); i += 1; continue
        if re.match(r"^\s*\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number"); p.paragraph_format.space_after = Pt(2)
            add_inline(p, re.sub(r"^\s*\d+\.\s+", "", line)); i += 1; continue
        if line.startswith(">"):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8)
            add_inline(p, line.lstrip("> "), base_italic=True); i += 1; continue
        if line.strip() in ("---", "***"):
            i += 1; continue
        if not line.strip():
            i += 1; continue
        # merge consecutive text lines into one paragraph
        j = i; buf = [line.strip()]
        while j + 1 < len(lines) and lines[j + 1].strip() and not re.match(r"^(#|\||```|\s*[-*]\s|\s*\d+\.\s|>|<<)", lines[j + 1]):
            j += 1; buf.append(lines[j].strip())
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, " ".join(buf))
        i = j + 1


# ----------------------------------------------------------------------------- result tables
def csv_table(name):
    p = TAB / f"paper_{name}.csv"
    return pd.read_csv(p) if p.exists() else None


def _fmt(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e6:
        return str(int(v))
    return v


def df_rows(df):
    return [list(df.columns)], [[_fmt(v) for v in r] for r in df.itertuples(index=False)]


def insert_result_table(spec):
    name, caption = spec.split("|", 1)
    df = None
    if name == "T1":
        df = csv_table("T1")
        df = df.rename(columns={"reg_domains": "registered domains", "https": "share https", "www": "share www."})
    elif name == "T2":
        df = csv_table("T2")
    elif name == "T10":
        df = csv_table("T10")
        if df is not None:
            df = df[["class", "feature_set", "drop_artifacts", "accuracy", "phiusiil", "lineage253k", "hannousse2021",
                     "ebbu2017", "phishstorm", "iscx2016"]].copy()
            df["representation"] = df.apply(lambda r: "lexical_v2 − artifact cols" if r.drop_artifacts else r.feature_set, axis=1)
            df = df[["class", "representation", "accuracy", "phiusiil", "lineage253k", "hannousse2021", "ebbu2017",
                     "phishstorm", "iscx2016"]].round(3)
    elif name == "EXT":
        t4 = csv_table("T4")
        if t4 is not None:
            d = t4[t4.run == "A_b28_grouped"].copy()
            cols = ["model", "in-dist AUC"] + [c for c in d.columns if c.endswith(" FPR") or c.endswith(" AUC")][:0]
            keep = ["model", "in-dist AUC", "hannousse2021 AUC", "hannousse2021 FPR", "ebbu2017 AUC", "ebbu2017 FPR",
                    "phishstorm AUC", "phishstorm FPR", "iscx2016 AUC", "iscx2016 FPR", "lineage253k AUC",
                    "lineage253k FPR", "phishtank recall"]
            df = d[[c for c in keep if c in d.columns]]
    elif name == "LOCO":
        t5 = csv_table("T5")
        if t5 is not None:
            df = t5[t5["model"].notna()] if "model" in t5.columns else t5
            df = df.drop(columns=[c for c in ("status",) if c in df.columns])
    if df is None or len(df) == 0:
        p = doc.add_paragraph(); add_inline(p, f"[{caption} — pending]", base_italic=True); return
    table_caption(caption)
    h, rows = df_rows(df)
    add_table(h[0], rows, PORTRAIT_TEXT_WIDTH, font=7.5)


# ----------------------------------------------------------------------------- manuscript preprocessing
ms = (ROOT / "paper/manuscript.md").read_text(encoding="utf-8")
title = re.search(r"^# (.+)$", ms, re.M).group(1)
abstract = re.search(r"## Abstract\s*\n(.*?)\n\*\*Keywords:\*\*", ms, re.S).group(1).strip()
abstract = re.sub(r"^\*Provisional[^\n]*\*\s*", "", abstract).strip()
keywords = re.search(r"\*\*Keywords:\*\*\s*(.+)", ms).group(1).strip()
body = ms[ms.index("## 1. Introduction"):ms.index("## References")]

swap = [
    ("```\nraw sources", "<<FIG|architecture.png|Pipeline of the study. All fitting uses training rows only and all model, threshold and calibration selection uses validation rows only; test, external corpora and feeds are scored once.>>\n```\nraw sources"),
    ("[Table T1 — generated; see `results/tables/paper_tables.md`.] ", "<<TABLE|T1|Corpora after cleaning (malicious = 1). Shares are over all URLs of the corpus.>>\n\n"),
    ("\n### 6.2 ", "\n<<TABLE|T2|Re-execution of the original pipeline code (random 60/20/20 split, seed 42, threshold 0.5), test split. The first four rows are the stored results of the original project on its unavailable dataset, shown for reference only.>>\n\n### 6.2 "),
    ("Figure 1 (`results/figures/dataset_artifacts.png`) and DATASET_AUDIT §3", "Fig. 2 and Appendix E"),
    ("(Table T10)", "(Table III)"),
    ("(Table T2)", "(Table II)"),
    ("\n### 6.3 ", "\n<<FIG|dataset_artifacts.png|Share of URLs with an https scheme, a leading www. label and an empty path (homepage), by class and corpus. Corpus B and the Kaggle corpus store URLs without a scheme.>>\n<<TABLE|T10|Dataset-identification probe: accuracy of a 6-way logistic regression predicting the source corpus of a URL from its representation (grouped 5-fold CV; chance = 0.167), overall and recall per corpus.>>\n\n### 6.3 "),
    ("(Table T3; `split_domain_overlap.png`)", "(Fig. 3; full results in Appendix G, Table T3)"),
    ("Despite this leakage, grouping hardly changes in-distribution performance on PhiUSIIL with the baseline features (5 seeds each):",
     "Despite this leakage, grouping hardly changes in-distribution performance on PhiUSIIL with the baseline features (5 seeds each; Table IV):\n\n<<MDTABLECAP|In-distribution test performance on corpus A with the original 28 features: random vs registered-domain-grouped split (mean ± sd over 5 seeds).>>"),
    ("\n### 6.4 ", "\n<<FIG|split_domain_overlap.png|Share of test URLs whose registered domain also occurs in the training partition, random vs grouped splits (mean ± sd over seeds).>>\n<<FIG|roc_pr_A_b28_grouped.png|ROC (log-scaled false-positive axis) and precision–recall curves on the corpus-A grouped test split, seed 42.>>\n\n### 6.4 "),
    ("(Figure 2, `external_roc_auc_A_b28_grouped.png`; Table T4)", "(Fig. 5; Table V; full results in Appendix G, Table T4)"),
    ("(FPR 0.994–1.000 on the four independent labelled corpora, 0.74–1.00 on corpus B; external AUC 0.34–0.78; Table T4)", "(FPR 0.994–1.000 on the four independent labelled corpora, 0.74–1.00 on corpus B; external AUC 0.34–0.78; Appendix G, Table T4)"),
    ("\n### 6.5 ", "\n<<FIG|external_roc_auc_A_b28_grouped.png|ROC-AUC on the corpus-A grouped test split and on the external labelled corpora for models trained on corpus A with the original features (mean ± sd over 5 seeds). Dotted line: chance.>>\n<<TABLE|EXT|External evaluation of corpus-A models (original features, grouped split, 5 seeds): ROC-AUC and the share of benign URLs flagged (FPR) at the threshold chosen on corpus-A validation data for FPR ≤ 1 %; PhishTank is positive-only (recall).>>\n<<FIG|reliability_A_b28_grouped.png|Reliability diagram on the corpus-A grouped test split (seed 42). Expected calibration error is at most 0.004 because almost all scores lie near 0 or 1; the few mid-range scores deviate visibly from the diagonal.>>\n\n### 6.5 "),
    ("\n### 6.6 ", "\n<<FIG|external_roc_auc_LOCO_lineage253k_v2.png|Leave-one-corpus-out with corpus B held out: ROC-AUC on the pool's own test split and on the held-out corpus (lexical_v2, 3 seeds).>>\n<<TABLE|LOCO|Leave-one-corpus-out results available so far: models trained on the other five corpora and tested on the held-out corpus (mean ± sd over 3 seeds). Recall and FPR use the threshold chosen on pool validation data for FPR ≤ 1 %.>>\n\n### 6.6 "),
    ("(T6)", "(Appendix G, T6)"), ("(T7, H2)", "(Appendix G, T7; H2)"), ("(T8)", "(Appendix G, T8)"), ("(T9;", "(Appendix G, T9;"),
]
for a, b in swap:
    if a not in body:
        print("WARNING: anchor not found:", a[:70])
    body = body.replace(a, b, 1)
# The text pipeline diagram in 4.1 is replaced by the figure.
body = re.sub(r"```\nraw sources.*?```\n", "", body, flags=re.S)
# Also replace the inline table in 6.4 "| Model | ROC-AUC random..." keeps caption via MDTABLECAP.

# ----------------------------------------------------------------------------- title page
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(60)
r = p.add_run(title); r.bold = True; r.font.size = Pt(20)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Research paper — full draft with literature survey, results and appendices"); r.italic = True; r.font.size = Pt(12)
for line in ("[Author names]", "[Department, Institution, City, Country]", "[Corresponding author e-mail]"):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(line); r.font.size = Pt(11)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(18)
p.add_run(f"Document generated {date.today().isoformat()} from the repository branch research-v2.").font.size = Pt(10)
status = doc.add_table(rows=1, cols=1); status.style = "Table Grid"; status.alignment = WD_TABLE_ALIGNMENT.CENTER
c = status.rows[0].cells[0]; c.width = Cm(14); shade(c, "FFF4D6")
cp = c.paragraphs[0]; rr = cp.add_run("Status of the results. "); rr.bold = True; rr.font.size = Pt(9.5)
rr = cp.add_run("Sections 1–5, the literature survey and all appendices are complete. Section 6 reports only experiment runs "
                "that have finished; every unfinished experiment is marked “pending” and no number is extrapolated. The original "
                "project's dataset (raw_data.csv) is not available; results use documented public corpora. Abstract, discussion "
                "and conclusion are provisional until the experiment matrix completes."); rr.font.size = Pt(9.5)
p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)

# abstract + keywords
h = doc.add_heading("Abstract", level=1)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; add_inline(p, abstract)
p = doc.add_paragraph(); r = p.add_run("Keywords — "); r.bold = True; r.italic = True; add_inline(p, keywords, base_italic=True)

# table of contents
cp_ = doc.add_paragraph(); cr_ = cp_.add_run("Contents"); cr_.bold = True; cr_.font.size = Pt(14)
add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
footer_page_numbers(doc.sections[0])

# ----------------------------------------------------------------------------- main text
handlers = {
    "FIG": lambda a: add_figure(FIG / a.split("|", 1)[0], a.split("|", 1)[1]),
    "TABLE": insert_result_table,
}
render_md(body, heading_offset=-1, handlers=handlers)

# references (computed after appendices so appendix citations are included) -> placeholder position
ref_anchor = doc.add_heading("References", level=1)
ref_par = doc.add_paragraph()  # placeholder; filled at the end

# ----------------------------------------------------------------------------- appendices
APP = [
    ("A", "Literature Survey Table", None),
    ("B", "Literature Review", "research/LITERATURE_REVIEW.md"),
    ("C", "Research Gaps, Hypotheses and Selected Direction", "research/RESEARCH_GAPS.md"),
    ("D", "Repository Audit of the Original Implementation", "research/REPOSITORY_AUDIT.md"),
    ("E", "Dataset Audit", "research/DATASET_AUDIT.md"),
    ("F", "Experiment Protocol", "research/EXPERIMENT_PROTOCOL.md"),
    ("G", "Complete Result Tables", None),
    ("H", "Additional Figures", None),
    ("I", "Decision Log", "research/DECISIONS.md"),
    ("J", "Reproducibility Guide", "README.md"),
]

# A: survey table (landscape)
new_section(landscape=True)
doc.add_heading("Appendix A. Literature Survey Table", level=1)
p = doc.add_paragraph()
add_inline(p, "One row per verified source (111 sources after removing three duplicate records of the same DOI). "
              "Numbers in the ‘Metrics reported’ column are the source authors' own and are not comparable across rows "
              "(datasets, splits, class ratios and metrics differ). Access: F = full text read, A = abstract or publisher page only, "
              "P = paywalled (metadata only). Full notes for every source are kept in the repository under research/reading_notes/.", size=9)
meta = {r["citekey"]: r for r in csv.DictReader(open(ROOT / "research/paper_metadata.csv", encoding="utf-8"))}
surv = [l for l in (ROOT / "research/LITERATURE_SURVEY_TABLE.md").read_text(encoding="utf-8").splitlines() if l.startswith("| ") and not l.startswith("| Citekey")]
rows = []
for l in surv:
    c = [x.strip() for x in l.strip().strip("|").split(" | ")]
    if len(c) < 12 or c[0] not in BIB:
        continue
    key = c[0]; m = meta.get(key, {})
    first = clean_tex(BIB[key].get("author", "")).split(" and ")[0].split(",")[0].split()[-1] if BIB[key].get("author") else key
    acc = {"full_text": "F", "abstract_only": "A", "paywalled": "P"}.get(m.get("access_status", ""), "?")
    rows.append([f"[{key}]", f"{first} ({m.get('year', '')}); {c[1].split(' (')[0]}", c[2], c[3], c[4], c[6], c[7], c[8], c[9], c[10], f"{c[11]} / {acc}"])
hdr = ["Ref.", "Author (year); venue", "Research objective", "Method", "Dataset", "Evaluation protocol", "Metrics reported",
       "Main findings", "Limitations", "Research gap", "Rel. / access"]
table_caption("Survey of verified literature (sorted by theme and year).", label="A1")
add_table(hdr, rows, LANDSCAPE_TEXT_WIDTH, font=6.2, col_weights=[4, 9, 12, 13, 11, 11, 13, 13, 13, 9, 4])

# B-F: markdown documents (portrait)
for letter, name, path in APP[1:6]:
    new_section(landscape=False)
    doc.add_heading(f"Appendix {letter}. {name}", level=1)
    render_md((ROOT / path).read_text(encoding="utf-8"), heading_offset=1, skip_title=True, table_font=7.5)

# G: complete result tables (landscape)
new_section(landscape=True)
doc.add_heading("Appendix G. Complete Result Tables", level=1)
p = doc.add_paragraph()
add_inline(p, "Generated by scripts/make_paper_tables.py from results/runs/*. Mean ± standard deviation over seeds. "
              "Rows marked pending correspond to experiments that had not finished when this document was generated.", size=9)
TITLES = {"T1": "Corpora after cleaning", "T2": "Original pipeline re-execution vs stored original results",
          "T3": "In-distribution test performance: random vs grouped splits",
          "T4": "External corpora: ROC-AUC and FPR at the validation-chosen FPR ≤ 1 % threshold",
          "T5": "Leave-one-corpus-out", "T6": "Paired comparisons against XGBoost", "T7": "Calibration and threshold transfer",
          "T8": "Sensitivity to label-preserving URL variants", "T9": "Computational cost",
          "T10": "Dataset-identification probe"}
for t in ["T1", "T2", "T3", "T4", "T5", "T6", "T7", "T8", "T9", "T10"]:
    df = csv_table(t)
    if df is None:
        table_caption(f"{TITLES[t]} — pending (analysis not yet run).", label=t)
        doc.add_paragraph()
        continue
    if t == "T9":
        df = df[~df.run_id.isin([])]
    table_caption(TITLES[t] + ".", label=t)
    if "status" in df.columns:
        df = df.copy()
        df["status"] = df["status"].where(df["status"].notna(), "complete")
        df = df.fillna("—")
    h, rws = df_rows(df)
    add_table(h[0], rws, LANDSCAPE_TEXT_WIDTH, font=6.5)

# H: additional figures (portrait)
new_section(landscape=False)
doc.add_heading("Appendix H. Additional Figures", level=1)
extra = [
    ("roc_pr_A_b28_random.png", "ROC and precision–recall curves, corpus A, random split, original features (seed 42)."),
    ("reliability_A_b28_random.png", "Reliability diagram, corpus A, random split (seed 42)."),
    ("external_roc_auc_A_b28_random.png", "In-distribution vs external ROC-AUC, corpus A, random split, original features (5 seeds)."),
    ("external_roc_auc_LOCO_phiusiil_v2.png", "Leave-one-corpus-out with PhiUSIIL held out (lexical_v2): pool test vs held-out ROC-AUC."),
    ("external_roc_auc_LOCO_phiusiil_v2noart.png", "Same as previous figure without the five artifact columns."),
    ("roc_pr_LOCO_lineage253k_v2.png", "ROC and precision–recall curves on the pool test split, corpus B held out (seed 42)."),
    ("temporal_recall_phishtank.png", "Preliminary: recall on PhishTank verified phishing URLs by submission year (bars: share of https URLs) at the validation-chosen FPR ≤ 1 % threshold. Years with fewer than 100 URLs omitted; analysis to be repeated on the complete matrix."),
    ("roc_curves.png", "Original project output (results/figures/roc_curves.png): ROC curves on the original, unavailable dataset; reproduced here unchanged for reference."),
    ("pr_curves.png", "Original project output: precision–recall curves on the original dataset."),
    ("model_comparison.png", "Original project output: model comparison on the original dataset (test split)."),
    ("feature_correlation_heatmap.png", "Original project output: correlation of the 15 highest-variance original features."),
]
for f, c in extra:
    add_figure(FIG / f, c, width_cm=14.5)
for f in ["logistic_regression", "random_forest", "xgboost", "weighted_soft-voting_ensemble"]:
    pth = ROOT / "results/confusion_matrices" / f"{f}_confusion_matrix.png"
    add_figure(pth, f"Original project output: confusion matrix of {f.replace('_', ' ')} on the original test split.", width_cm=8.5)
add_figure(ROOT / "results/shap/shap_summary_ensemble.png", "Original project output: SHAP global importance of the soft-voting ensemble (KernelExplainer, 500 validation URLs). https_usage and num_subdomains dominate.", width_cm=13)

# I, J
for letter, name, path in APP[8:]:
    new_section(landscape=False)
    doc.add_heading(f"Appendix {letter}. {name}", level=1)
    txt = (ROOT / path).read_text(encoding="utf-8")
    if path == "README.md":
        txt = txt.replace("](", "] (")  # keep link text readable
    render_md(txt, heading_offset=1, skip_title=True, table_font=7.0)

# ----------------------------------------------------------------------------- fill references
# references are inserted at the placeholder paragraph, in citation order
uncited = [k for k in BIB if k not in CITE_ORDER]
for k in uncited:
    cite_num(k)
anchor = ref_par
intro = anchor.insert_paragraph_before()
add_inline(intro, f"Numbered in order of first citation. References {len(CITE_ORDER) - len(uncited) + 1}–{len(CITE_ORDER)} "
                  "are verified sources reviewed for the literature survey but not cited in the text." if uncited else
                  "Numbered in order of first citation. All entries were verified against an authoritative page "
                  "(DOI resolver, publisher, arXiv or proceedings site); see research/paper_metadata.csv.", size=9)
for n, key in enumerate(CITE_ORDER, 1):
    q = anchor.insert_paragraph_before()
    q.paragraph_format.left_indent = Cm(0.9); q.paragraph_format.first_line_indent = Cm(-0.9)
    q.paragraph_format.space_after = Pt(2)
    r = q.add_run(f"[{n}] "); r.font.size = Pt(9)
    r = q.add_run(fmt_reference(BIB[key])); r.font.size = Pt(9)
anchor._element.getparent().remove(anchor._element)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("saved", OUT, "| figures:", FIG_N[0], "| captioned tables:", TAB_N[0], "| references:", len(CITE_ORDER))
